"""Build the interactive Gallotta family tree from the Excel workbook.

Reads the "Persone" and "Relazioni" sheets of Albero_Genealogico_Gallotta.xlsx,
checks the data for common mistakes and writes docs/data.js, which is loaded by
the viewer in docs/index.html.

Usage:
    python genealogic.py            # rebuild docs/data.js
    python genealogic.py --open     # rebuild and open the tree in the browser
"""

import argparse
import json
import webbrowser
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
EXCEL = ROOT / "Albero_Genealogico_Gallotta.xlsx"
OUT = ROOT / "docs" / "data.js"


def clean(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def read_sheet(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    records = []
    for row in rows:
        if all(v is None for v in row):
            continue
        records.append({h: clean(v) for h, v in zip(header, row) if h})
    return records


def load_data(path=EXCEL):
    wb = load_workbook(path, read_only=True, data_only=True)
    people = read_sheet(wb, "Persone")
    relations = read_sheet(wb, "Relazioni")

    persons = [
        {
            "id": p["person_id"],
            "nome": p.get("nome") or "",
            "cognome": p.get("cognome") or "",
            "soprannome": p.get("soprannome"),
            "gen": p.get("generazione"),
        }
        for p in people
        if p.get("person_id") is not None
    ]
    rels = [
        {"tipo": r["tipo"], "a": r["person_id_1"], "b": r["person_id_2"]}
        for r in relations
        if r.get("tipo") and r.get("person_id_1") is not None and r.get("person_id_2") is not None
    ]
    return persons, rels


def check(persons, rels):
    """Return a list of human readable warnings about the data."""
    warnings = []
    by_id = {}
    for p in persons:
        if p["id"] in by_id:
            warnings.append(f"person_id {p['id']} is used more than once")
        by_id[p["id"]] = p

    def name(pid):
        p = by_id.get(pid)
        return f"{p['nome']} {p['cognome']} ({pid})" if p else f"#{pid}"

    parents = defaultdict(set)
    seen = set()
    for r in rels:
        key = (r["tipo"], r["a"], r["b"])
        if key in seen:
            warnings.append(f"duplicate relation: {r['tipo']} {name(r['a'])} -> {name(r['b'])}")
        seen.add(key)
        for pid in (r["a"], r["b"]):
            if pid not in by_id:
                warnings.append(f"relation '{r['tipo']}' refers to unknown person_id {pid}")
        if r["tipo"] == "genitore":
            parents[r["b"]].add(r["a"])
        elif r["tipo"] != "matrimonio":
            warnings.append(f"unknown relation type '{r['tipo']}' ({name(r['a'])}, {name(r['b'])})")

    married = {r["a"] for r in rels if r["tipo"] == "matrimonio"}
    married |= {r["b"] for r in rels if r["tipo"] == "matrimonio"}
    min_gen = min((p["gen"] for p in persons if p["gen"] is not None), default=1)

    for p in persons:
        pid = p["id"]
        ps = parents.get(pid, set())
        if len(ps) > 2:
            warnings.append(f"{name(pid)} has more than two parents")
        for par in ps:
            g_par = by_id.get(par, {}).get("gen")
            if p["gen"] is not None and g_par is not None and p["gen"] != g_par + 1:
                warnings.append(
                    f"{name(pid)} is generation {p['gen']} but parent {name(par)} is generation {g_par}"
                )
        if not ps and p["gen"] is not None and p["gen"] > min_gen and pid not in married:
            warnings.append(f"{name(pid)} has no parents and no spouse: not connected to the tree")
    return warnings


def unlinked_roots(persons, rels):
    """People below the first generation with no parents whose spouse has none either.

    These start a separate branch, usually because their parents are missing
    from the "Relazioni" sheet.
    """
    parents = defaultdict(set)
    spouses = defaultdict(set)
    for r in rels:
        if r["tipo"] == "genitore":
            parents[r["b"]].add(r["a"])
        elif r["tipo"] == "matrimonio":
            spouses[r["a"]].add(r["b"])
            spouses[r["b"]].add(r["a"])
    min_gen = min((p["gen"] for p in persons if p["gen"] is not None), default=1)
    roots = []
    for p in persons:
        if p["gen"] is None or p["gen"] <= min_gen or parents[p["id"]]:
            continue
        if any(parents[s] for s in spouses[p["id"]]):
            continue
        roots.append(p)
    return roots


def build(open_browser=False):
    persons, rels = load_data()
    warnings = check(persons, rels)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    def rows(items):
        return ",\n".join("    " + json.dumps(i, ensure_ascii=False) for i in items)

    OUT.write_text(
        "// Generated by genealogic.py from Albero_Genealogico_Gallotta.xlsx. Do not edit by hand.\n"
        "window.FAMILY_DATA = {\n"
        '  "title": "Albero Genealogico Gallotta",\n'
        f'  "persons": [\n{rows(persons)}\n  ],\n'
        f'  "relations": [\n{rows(rels)}\n  ]\n'
        "};\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT.relative_to(ROOT)}: {len(persons)} people, {len(rels)} relations")

    roots = unlinked_roots(persons, rels)
    if roots:
        names = ", ".join(f"{p['nome']} {p['cognome']} ({p['id']})" for p in roots)
        print(f"Note: these people have no parents in the data and start their own branch: {names}")
    for w in warnings:
        print(f"Warning: {w}")

    if open_browser:
        webbrowser.open((ROOT / "docs" / "index.html").as_uri())


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--open", action="store_true", help="open the tree in the browser after building")
    build(open_browser=parser.parse_args().open)
